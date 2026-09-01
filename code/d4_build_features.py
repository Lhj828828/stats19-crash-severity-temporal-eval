"""Create the frozen D4 collision-level feature table.

Raw annual CSV files are read-only inputs. The output keeps metadata, target,
and model features visibly separated by column prefixes. No rows are dropped,
no values are statistically imputed, and no encoder is fitted at D4.
"""

from __future__ import annotations

import csv
import gzip
import hashlib
import io
import json
import os
import re
from collections import Counter
from contextlib import contextmanager
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parents[1]
RAW_DIR = PROJECT_DIR / "data" / "raw" / "collisions"
INTERIM_DIR = PROJECT_DIR / "data" / "interim"
DOCUMENT_DIR = PROJECT_DIR / "data" / "external" / "documentation"
LOG_DIR = PROJECT_DIR / "logs"
CONFIG_DIR = PROJECT_DIR / "config"

MANIFEST_FILE = CONFIG_DIR / "d3_feature_manifest.json"
GUIDE_FILE = (
    DOCUMENT_DIR
    / "dft-road-casualty-statistics-road-safety-open-dataset-data-guide-2025.xlsx"
)
GUIDE_SHEET = "2024_code_list"
OUTPUT_FILE = INTERIM_DIR / "d4_modeling_table.csv.gz"
YEARS = tuple(range(2018, 2025))
VALID_SPEED_LIMITS = {20, 30, 40, 50, 60, 70}
UNKNOWN_PATTERN = re.compile(
    r"unknown|missing|out of range|not reported|not applicable", re.I
)

META_COLUMNS = [
    "meta_collision_index",
    "meta_collision_year",
    "meta_date",
    "meta_police_force",
    "meta_injury_based",
]
TARGET_COLUMN = "target_severity"


def normalize_code(value: object) -> str | None:
    if value is None or pd.isna(value):
        return None
    text = str(value).strip()
    if re.fullmatch(r"-?\d+(?:\.0+)?", text):
        return str(int(float(text)))
    return text


def load_manifest() -> dict[str, object]:
    manifest = json.loads(MANIFEST_FILE.read_text(encoding="utf-8"))
    if manifest.get("status") != "FROZEN":
        raise ValueError("D3 feature manifest is not frozen")
    return manifest


def load_code_dictionary() -> dict[str, dict[str, str]]:
    workbook = load_workbook(GUIDE_FILE, read_only=True, data_only=True)
    if GUIDE_SHEET not in workbook.sheetnames:
        raise ValueError(f"Official guide worksheet not found: {GUIDE_SHEET}")
    sheet = workbook[GUIDE_SHEET]
    mappings: dict[str, dict[str, str]] = {}
    for table, field, code, label, _note in sheet.iter_rows(values_only=True):
        if table != "collision" or not field or not label:
            continue
        normalized = normalize_code(code)
        if normalized is None:
            continue
        mappings.setdefault(str(field), {})[normalized] = str(label).strip()
    workbook.close()
    return mappings


def encode_categorical(
    values: pd.Series,
    field: str,
    mapping: dict[str, str],
) -> tuple[pd.Series, pd.Series]:
    normalized = values.astype("string").str.strip().str.replace(
        r"\.0+$", "", regex=True
    )
    observed = set(normalized.dropna().unique().tolist())
    unexpected = sorted(observed - set(mapping))
    if unexpected:
        raise ValueError(
            f"{field} contains codes absent from the official guide: {unexpected}"
        )
    encoded = normalized.map(
        {code: f"{code}:{label}" for code, label in mapping.items()}
    ).fillna("__CSV_MISSING__")
    return encoded, normalized


def hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(8 * 1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


@contextmanager
def deterministic_gzip_text(path: Path):
    with path.open("wb") as binary_output:
        with gzip.GzipFile(
            filename="",
            mode="wb",
            fileobj=binary_output,
            compresslevel=6,
            mtime=0,
        ) as compressed_output:
            with io.TextIOWrapper(
                compressed_output, encoding="utf-8", newline=""
            ) as text_output:
                yield text_output


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    if not rows:
        raise ValueError(f"No rows generated for {path.name}")
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    manifest = load_manifest()
    code_dictionary = load_code_dictionary()
    categorical_fields = list(manifest["direct_features"]["categorical"])
    numeric_fields = list(manifest["direct_features"]["numeric"])
    if numeric_fields != ["speed_limit"]:
        raise ValueError(f"Unexpected D3 numeric feature list: {numeric_fields}")

    feature_columns = [f"feature_{field}" for field in categorical_fields]
    feature_columns += ["feature_speed_limit", "feature_month", "feature_hour"]
    expected_output_columns = META_COLUMNS + [TARGET_COLUMN] + feature_columns
    if len(feature_columns) != 17 or len(set(feature_columns)) != 17:
        raise AssertionError("Expected 17 unique D4 feature columns")

    for field in categorical_fields:
        if field not in code_dictionary:
            raise ValueError(f"No official categorical mapping found for {field}")
    speed_mapping = code_dictionary.get("speed_limit", {})
    guide_speed_unknown = {
        int(code)
        for code, label in speed_mapping.items()
        if re.fullmatch(r"-?\d+", code) and UNKNOWN_PATTERN.search(label)
    }

    input_columns = list(
        dict.fromkeys(
            [
                "collision_index",
                "collision_year",
                "date",
                "time",
                "police_force",
                "collision_injury_based",
                "collision_severity",
                *categorical_fields,
                "speed_limit",
            ]
        )
    )

    INTERIM_DIR.mkdir(parents=True, exist_ok=True)
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    part_file = OUTPUT_FILE.with_suffix(OUTPUT_FILE.suffix + ".part")
    part_file.unlink(missing_ok=True)
    year_audit: list[dict[str, object]] = []
    category_audit: list[dict[str, object]] = []
    total_rows = 0
    observed_speed_unknown: set[int] = set()

    try:
        with deterministic_gzip_text(part_file) as output_handle:
            write_header = True
            for year in YEARS:
                source = RAW_DIR / f"collision_{year}.csv"
                raw = pd.read_csv(
                    source,
                    usecols=input_columns,
                    dtype=str,
                    keep_default_na=True,
                    low_memory=False,
                )
                rows = len(raw)
                if rows == 0:
                    raise ValueError(f"No records found for {year}")
                parsed_year = pd.to_numeric(raw["collision_year"], errors="raise")
                if not parsed_year.eq(year).all():
                    raise ValueError(f"collision_year mismatch in {source.name}")

                parsed_date = pd.to_datetime(
                    raw["date"], format="%d/%m/%Y", errors="coerce"
                )
                invalid_dates = int(parsed_date.isna().sum())
                time_text = raw["time"].astype("string").str.strip()
                valid_time = time_text.str.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d")
                invalid_times = int((~valid_time.fillna(False)).sum())
                if invalid_dates or invalid_times:
                    raise ValueError(
                        f"{year}: invalid dates={invalid_dates}, invalid times={invalid_times}"
                    )

                target_raw = raw["collision_severity"].map(normalize_code)
                target_map = {"3": 0, "2": 1, "1": 2}
                unexpected_target = sorted(
                    set(target_raw.dropna().unique().tolist()) - set(target_map)
                )
                if unexpected_target or target_raw.isna().any():
                    raise ValueError(
                        f"{year}: invalid target codes={unexpected_target}, "
                        f"missing={int(target_raw.isna().sum())}"
                    )

                output = pd.DataFrame(index=raw.index)
                output["meta_collision_index"] = raw["collision_index"]
                output["meta_collision_year"] = parsed_year.astype("int16")
                output["meta_date"] = parsed_date.dt.strftime("%Y-%m-%d")
                output["meta_police_force"] = raw["police_force"].map(normalize_code)
                output["meta_injury_based"] = raw["collision_injury_based"].map(
                    normalize_code
                )
                output[TARGET_COLUMN] = target_raw.map(target_map).astype("int8")

                for field in categorical_fields:
                    encoded, normalized = encode_categorical(
                        raw[field], field, code_dictionary[field]
                    )
                    output[f"feature_{field}"] = encoded
                    counts = normalized.fillna("__CSV_MISSING__").value_counts(
                        dropna=False
                    )
                    for code, count in counts.items():
                        label = (
                            "CSV missing"
                            if code == "__CSV_MISSING__"
                            else code_dictionary[field][str(code)]
                        )
                        category_audit.append(
                            {
                                "year": year,
                                "field_name": field,
                                "raw_code": code,
                                "official_label": label,
                                "count": int(count),
                                "pct_of_year": int(count) / rows,
                                "handling": "explicit_unknown_or_not_applicable_category"
                                if UNKNOWN_PATTERN.search(label)
                                else "explicit_category",
                            }
                        )

                speed = pd.to_numeric(raw["speed_limit"], errors="raise")
                observed_speed = {int(value) for value in speed.dropna().unique()}
                unknown_speed = observed_speed - VALID_SPEED_LIMITS
                unverified_speed = unknown_speed - guide_speed_unknown
                if unverified_speed:
                    raise ValueError(
                        f"{year}: speed_limit contains unverified codes "
                        f"{sorted(unverified_speed)}"
                    )
                observed_speed_unknown.update(unknown_speed)
                output["feature_speed_limit"] = speed.mask(
                    speed.isin(guide_speed_unknown), np.nan
                ).astype("float32")
                output["feature_month"] = "M" + parsed_date.dt.month.astype(
                    "string"
                ).str.zfill(2)
                output["feature_hour"] = "H" + time_text.str.slice(0, 2)

                if output.columns.tolist() != expected_output_columns:
                    raise AssertionError(
                        f"Unexpected output schema for {year}: {output.columns.tolist()}"
                    )
                if output["meta_collision_index"].isna().any():
                    raise ValueError(f"{year}: missing collision identifiers")
                feature_missing = {
                    column: int(output[column].isna().sum())
                    for column in feature_columns
                }
                unexpected_missing = {
                    field: count
                    for field, count in feature_missing.items()
                    if count and field != "feature_speed_limit"
                }
                if unexpected_missing:
                    raise ValueError(
                        f"{year}: unexpected missing feature values: {unexpected_missing}"
                    )

                output.to_csv(
                    output_handle,
                    index=False,
                    header=write_header,
                    lineterminator="\n",
                    na_rep="",
                )
                write_header = False
                total_rows += rows
                target_counts = Counter(output[TARGET_COLUMN].tolist())
                year_audit.append(
                    {
                        "year": year,
                        "input_rows": rows,
                        "output_rows": len(output),
                        "slight_count": target_counts[0],
                        "serious_count": target_counts[1],
                        "fatal_count": target_counts[2],
                        "speed_missing_after_mapping": int(
                            output["feature_speed_limit"].isna().sum()
                        ),
                        "invalid_date_count": invalid_dates,
                        "invalid_time_count": invalid_times,
                    }
                )
        os.replace(part_file, OUTPUT_FILE)
    except Exception:
        part_file.unlink(missing_ok=True)
        raise

    output_sha256 = hash_file(OUTPUT_FILE)
    write_csv(LOG_DIR / "d4_year_audit.csv", year_audit)
    write_csv(LOG_DIR / "d4_category_audit.csv", category_audit)
    output_schema = {
        "version": "D4_V1",
        "status": "GENERATED_NOT_D5_CLEANED",
        "file": OUTPUT_FILE.relative_to(PROJECT_DIR).as_posix(),
        "compression": "gzip",
        "sha256": output_sha256,
        "row_count": total_rows,
        "metadata_columns": META_COLUMNS,
        "metadata_roles": {
            "meta_collision_index": "record trace and audit only",
            "meta_collision_year": "temporal split only",
            "meta_date": "chronological ordering only",
            "meta_police_force": "regional grouping and sensitivity analysis only",
            "meta_injury_based": "reporting-system sensitivity analysis only",
        },
        "target_column": TARGET_COLUMN,
        "feature_prefix": "feature_",
        "feature_columns": feature_columns,
        "feature_selection_policy": "Model code must select exactly feature_columns from this schema; metadata and target columns are never selected by exclusion logic.",
        "categorical_feature_columns": [
            f"feature_{field}" for field in categorical_fields
        ]
        + ["feature_month", "feature_hour"],
        "numeric_feature_columns": ["feature_speed_limit"],
        "guide_speed_unknown_codes": sorted(guide_speed_unknown),
        "observed_speed_unknown_codes": sorted(observed_speed_unknown),
        "missing_value_policy": {
            "LightGBM": "native missing-value handling",
            "multinomial_logistic": "imputation fitted within each training fold only",
            "ordered_logit": "imputation fitted within each training fold only",
        },
        "row_policy": "No rows dropped at D4",
    }
    (CONFIG_DIR / "d4_output_schema.json").write_text(
        json.dumps(output_schema, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    total_target = Counter()
    total_speed_missing = 0
    for row in year_audit:
        total_target[0] += int(row["slight_count"])
        total_target[1] += int(row["serious_count"])
        total_target[2] += int(row["fatal_count"])
        total_speed_missing += int(row["speed_missing_after_mapping"])
    speed_missing_pct = total_speed_missing / total_rows
    lines = [
        "# D4 feature-engineering audit",
        "",
        f"- Input/output rows: **{total_rows:,}**; no rows dropped.",
        f"- Model features before encoding: **{len(feature_columns)}**.",
        f"- Metadata columns: **{len(META_COLUMNS)}**; all use the `meta_` prefix.",
        f"- Target column: `{TARGET_COLUMN}` with 0=Slight, 1=Serious, 2=Fatal.",
        f"- Target counts: Slight={total_target[0]:,}, Serious={total_target[1]:,}, Fatal={total_target[2]:,}.",
        f"- Speed-limit valid values: {sorted(VALID_SPEED_LIMITS)}.",
        f"- Speed-limit unknown codes listed by the current guide: {sorted(guide_speed_unknown)}.",
        f"- Speed-limit unknown codes observed in 2018-2024: {sorted(observed_speed_unknown)}; mapped to missing rows={total_speed_missing:,} ({speed_missing_pct:.3%}).",
        "- LightGBM uses native missing-value handling; multinomial logistic regression and ordered logit use imputation fitted within each training fold only.",
        "- Future model code must select exactly the 17-column feature_columns allowlist in d4_output_schema.json; metadata are never admitted through blacklist-style exclusion.",
        "- Categorical codes were mapped to `code:official label`; unknown and not-applicable labels remain explicit categories.",
        "- No one-hot encoding, statistical imputation, resampling or scaling was performed at D4.",
        f"- Output SHA-256: `{output_sha256}`.",
        "",
        "The D4 table is an interim artifact. D5 must still perform duplicate, missingness and sample-flow quality control before a final processed dataset is frozen.",
    ]
    (LOG_DIR / "d4_summary.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

    print("D4 feature engineering completed.")
    print(f"  rows: {total_rows:,}")
    print(f"  features: {len(feature_columns)}")
    print(f"  output: {OUTPUT_FILE}")
    print(f"  sha256: {output_sha256}")


if __name__ == "__main__":
    main()
