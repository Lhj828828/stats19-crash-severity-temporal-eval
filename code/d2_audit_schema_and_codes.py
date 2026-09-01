"""Audit STATS19 collision schema and category codes for 2018-2024.

The script reads, but never modifies, the annual raw CSV files and the official
DfT data guide. It writes reproducible D2 audit tables to the project log
directory. Feature inclusion and leakage decisions are deliberately deferred
to D3.
"""

from __future__ import annotations

import csv
import re
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parents[1]
RAW_DIR = PROJECT_DIR / "data" / "raw" / "collisions"
DOCUMENT_DIR = PROJECT_DIR / "data" / "external" / "documentation"
LOG_DIR = PROJECT_DIR / "logs"

GUIDE_FILE = (
    DOCUMENT_DIR
    / "dft-road-casualty-statistics-road-safety-open-dataset-data-guide-2025.xlsx"
)
GUIDE_SHEET = "2024_code_list"
YEARS = tuple(range(2018, 2025))
CHUNK_SIZE = 50_000

TARGET_FIELD = "collision_severity"
REVIEWED_CANDIDATES = {
    "day_of_week",
    "time",
    "road_type",
    "speed_limit",
    "junction_detail",
    "junction_control",
    "light_conditions",
    "weather_conditions",
    "road_surface_conditions",
    "urban_or_rural_area",
}
IDENTIFIER_OR_LOCATION_FIELDS = {
    "collision_index",
    "collision_ref_no",
    "location_easting_osgr",
    "location_northing_osgr",
    "longitude",
    "latitude",
    "local_authority_district",
    "local_authority_ons_district",
    "local_authority_highway",
    "local_authority_highway_current",
    "lsoa_of_accident_location",
}
HISTORIC_FIELDS = {
    "junction_detail_historic",
    "pedestrian_crossing_human_control_historic",
    "pedestrian_crossing_physical_facilities_historic",
    "carriageway_hazards_historic",
}
CURRENT_REPLACEMENT_FIELDS = {
    "junction_detail",
    "pedestrian_crossing",
    "carriageway_hazards",
}
NON_CATEGORICAL_CODELIKE_FIELDS = {
    "first_road_number",
    "second_road_number",
    "speed_limit",
}
OUTCOME_RELATED_FIELDS = {
    "number_of_casualties",
    "enhanced_severity_collision",
    "collision_injury_based",
    "collision_adjusted_severity_serious",
    "collision_adjusted_severity_slight",
}
OFFICIAL_FIELD_ALIASES = {
    "enhanced_severity_collision": "enhanced_collision_severity",
    "collision_adjusted_severity_serious": "collision_adjusted_serious",
    "collision_adjusted_severity_slight": "collision_adjusted_slight",
}


def clean_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_code(value: object) -> str | None:
    """Normalize discrete integer codes; reject ranges and free-form formats."""
    text = clean_cell(value)
    if not text:
        return None
    if re.fullmatch(r"-?\d+(?:\.0+)?", text):
        return str(int(float(text)))
    return None


def official_field_name(field: str) -> str:
    return OFFICIAL_FIELD_ALIASES.get(field, field)


def dictionary_entries(
    field: str, dictionary: dict[str, list[dict[str, str]]]
) -> list[dict[str, str]]:
    return dictionary.get(official_field_name(field), [])


def read_official_dictionary() -> tuple[dict[str, list[dict[str, str]]], list[str]]:
    if not GUIDE_FILE.exists():
        raise FileNotFoundError(f"Official DfT guide not found: {GUIDE_FILE}")

    workbook = load_workbook(GUIDE_FILE, read_only=True, data_only=True)
    if GUIDE_SHEET not in workbook.sheetnames:
        raise ValueError(f"Expected worksheet not found: {GUIDE_SHEET}")
    sheet = workbook[GUIDE_SHEET]

    rows = sheet.iter_rows(values_only=True)
    header = [clean_cell(value) for value in next(rows)]
    expected_header = ["table", "field name", "code/format", "label", "note"]
    if header[:5] != expected_header:
        raise ValueError(f"Unexpected guide header: {header[:5]}")

    dictionary: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        table, field, code, label, note = [clean_cell(value) for value in row[:5]]
        if table != "collision" or not field:
            continue
        dictionary[field].append(
            {"code_or_format": code, "label": label, "note": note}
        )
    workbook.close()
    return dict(dictionary), header


def field_role(field: str) -> str:
    if field == TARGET_FIELD:
        return "target"
    if field in REVIEWED_CANDIDATES:
        return "reviewed_candidate"
    if field in OUTCOME_RELATED_FIELDS:
        return "outcome_related_D3_review"
    if field in HISTORIC_FIELDS:
        return "historic_version"
    if field in CURRENT_REPLACEMENT_FIELDS:
        return "current_harmonized_version"
    if field in IDENTIFIER_OR_LOCATION_FIELDS:
        return "identifier_or_location_D3_review"
    return "not_yet_decided"


def get_headers() -> dict[int, list[str]]:
    headers: dict[int, list[str]] = {}
    for year in YEARS:
        path = RAW_DIR / f"collision_{year}.csv"
        if not path.exists():
            raise FileNotFoundError(f"Annual file not found: {path}")
        headers[year] = pd.read_csv(path, nrows=0).columns.tolist()
    return headers


def audit_data(
    headers: dict[int, list[str]], dictionary: dict[str, list[dict[str, str]]]
) -> tuple[list[dict[str, object]], list[dict[str, object]], dict[int, int]]:
    all_fields = list(dict.fromkeys(field for year in YEARS for field in headers[year]))
    coded_fields = {
        field
        for field in all_fields
        if (
            entries := dictionary_entries(field, dictionary)
        )
        and
        sum(
            normalize_code(entry["code_or_format"]) is not None
            and bool(entry["label"])
            for entry in entries
        )
        >= 2
    } - NON_CATEGORICAL_CODELIKE_FIELDS
    audited_fields = sorted(coded_fields.intersection(all_fields))
    counted_fields = sorted(
        (set(audited_fields) | REVIEWED_CANDIDATES | {TARGET_FIELD}).intersection(
            all_fields
        )
    )

    schema_rows: list[dict[str, object]] = []
    code_rows: list[dict[str, object]] = []
    row_counts: dict[int, int] = {}

    official_lookup: dict[str, dict[str, str]] = {}
    unknown_code_lookup: dict[str, set[str]] = {}
    for field in audited_fields:
        official_lookup[field] = {
            code: entry["label"]
            for entry in dictionary_entries(field, dictionary)
            if (code := normalize_code(entry["code_or_format"])) is not None
        }
    for field in counted_fields:
        unknown_code_lookup[field] = {
            code
            for entry in dictionary_entries(field, dictionary)
            if (code := normalize_code(entry["code_or_format"])) is not None
            and re.search(
                r"unknown|missing|out of range|not reported", entry["label"], re.I
            )
        }

    for year in YEARS:
        path = RAW_DIR / f"collision_{year}.csv"
        nonmissing = Counter()
        value_counts: dict[str, Counter[str]] = {
            field: Counter() for field in counted_fields if field in headers[year]
        }
        total_rows = 0

        for chunk in pd.read_csv(
            path,
            dtype=str,
            keep_default_na=True,
            chunksize=CHUNK_SIZE,
            low_memory=False,
        ):
            total_rows += len(chunk)
            for field in headers[year]:
                nonmissing[field] += int(chunk[field].notna().sum())
            for field in value_counts:
                values = chunk[field].dropna().astype(str).str.strip()
                value_counts[field].update(values[values.ne("")].tolist())

        row_counts[year] = total_rows
        for field in all_fields:
            present = field in headers[year]
            count = int(nonmissing[field]) if present else 0
            unknown_count = (
                sum(
                    value_counts[field].get(code, 0)
                    for code in unknown_code_lookup.get(field, set())
                )
                if field in value_counts
                else 0
            )
            format_issue_count = 0
            if field == "time" and field in value_counts:
                format_issue_count = sum(
                    value_count
                    for value, value_count in value_counts[field].items()
                    if re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", value) is None
                )
            elif field == "speed_limit" and field in value_counts:
                allowed_speed_values = {"-1", "20", "30", "40", "50", "60", "70"}
                format_issue_count = sum(
                    value_count
                    for value, value_count in value_counts[field].items()
                    if value not in allowed_speed_values
                )
            schema_rows.append(
                {
                    "year": year,
                    "field_name": field,
                    "present": present,
                    "column_position": headers[year].index(field) + 1 if present else "",
                    "row_count": total_rows,
                    "nonmissing_count": count,
                    "nonmissing_pct": count / total_rows if total_rows else 0.0,
                    "csv_blank_count": total_rows - count if present else total_rows,
                    "csv_blank_pct": (total_rows - count) / total_rows
                    if total_rows
                    else 0.0,
                    "official_unknown_code_count": unknown_count,
                    "official_unknown_code_pct": unknown_count / total_rows
                    if total_rows
                    else 0.0,
                    "format_issue_count": format_issue_count,
                    "format_issue_pct": format_issue_count / total_rows
                    if total_rows
                    else 0.0,
                    "official_field_name": official_field_name(field),
                    "dictionary_name_match": "alias"
                    if field in OFFICIAL_FIELD_ALIASES
                    else "exact",
                    "official_dictionary_entry": bool(
                        dictionary_entries(field, dictionary)
                    ),
                    "d2_role": field_role(field),
                }
            )

        for field in audited_fields:
            if field not in value_counts:
                continue
            counts = value_counts[field]
            expected = official_lookup[field]
            for code, count in sorted(
                counts.items(), key=lambda item: (not item[0].lstrip("-").isdigit(), item[0])
            ):
                normalized = normalize_code(code)
                label = expected.get(normalized or "", "")
                code_rows.append(
                    {
                        "year": year,
                        "field_name": field,
                        "observed_code": code,
                        "official_label": label,
                        "count": count,
                        "pct_of_year": count / total_rows if total_rows else 0.0,
                        "listed_in_official_guide": bool(label),
                        "unknown_or_missing_label": bool(
                            re.search(r"unknown|missing|out of range|not reported", label, re.I)
                        ),
                    }
                )

    return schema_rows, code_rows, row_counts


def build_mapping(
    headers: dict[int, list[str]], dictionary: dict[str, list[dict[str, str]]]
) -> list[dict[str, object]]:
    reference_fields = headers[YEARS[0]]
    rows: list[dict[str, object]] = []
    for position, field in enumerate(reference_fields, start=1):
        official_name = official_field_name(field)
        entries = dictionary_entries(field, dictionary)
        formats = sorted(
            {
                entry["code_or_format"]
                for entry in entries
                if entry["code_or_format"]
                and normalize_code(entry["code_or_format"]) is None
            }
        )
        coded_labels = [
            f"{normalize_code(entry['code_or_format'])}={entry['label']}"
            for entry in entries
            if normalize_code(entry["code_or_format"]) is not None and entry["label"]
        ]
        notes = sorted({entry["note"] for entry in entries if entry["note"]})
        years_present = [str(year) for year in YEARS if field in headers[year]]
        rows.append(
            {
                "column_position": position,
                "field_name": field,
                "official_field_name": official_name,
                "dictionary_name_match": "alias"
                if field in OFFICIAL_FIELD_ALIASES
                else "exact",
                "official_dictionary_entry": bool(entries),
                "years_present": ";".join(years_present),
                "present_all_years": len(years_present) == len(YEARS),
                "official_format": "; ".join(formats),
                "official_codes_and_labels": " | ".join(coded_labels),
                "official_note": " | ".join(notes),
                "d2_role": field_role(field),
                "final_D3_decision": "PENDING",
            }
        )
    return rows


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    if not rows:
        raise ValueError(f"No rows generated for {path.name}")
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def write_summary(
    headers: dict[int, list[str]],
    schema_rows: list[dict[str, object]],
    code_rows: list[dict[str, object]],
    row_counts: dict[int, int],
) -> None:
    reference = headers[YEARS[0]]
    schema_identical = all(headers[year] == reference for year in YEARS)
    absent = [row for row in schema_rows if not row["present"]]
    unexpected = [row for row in code_rows if not row["listed_in_official_guide"]]
    format_issue_count = sum(int(row["format_issue_count"]) for row in schema_rows)
    first_year_rows = [row for row in schema_rows if int(row["year"]) == YEARS[0]]
    dictionary_missing = [
        row["field_name"]
        for row in first_year_rows
        if not row["official_dictionary_entry"]
    ]
    alias_fields = [
        row["field_name"]
        for row in first_year_rows
        if row["dictionary_name_match"] == "alias"
    ]
    severity_rows = [
        row for row in code_rows if row["field_name"] == TARGET_FIELD
    ]
    severity_codes = sorted({str(row["observed_code"]) for row in severity_rows})
    severity_expected = severity_codes == ["1", "2", "3"]

    lines = [
        "# D2 schema and code audit",
        "",
        f"- Audit scope: STATS19 collision files, {YEARS[0]}-{YEARS[-1]}.",
        f"- Official dictionary: `{GUIDE_FILE.name}`, worksheet `{GUIDE_SHEET}`.",
        f"- Annual schemas identical: **{schema_identical}**.",
        f"- Columns in each annual file: **{len(reference)}**.",
        f"- Fields absent in any year: **{len(absent)}**.",
        f"- Raw fields requiring a documented official-name alias: **{len(alias_fields)}**.",
        f"- Raw fields with no official dictionary match after aliasing: **{len(dictionary_missing)}**.",
        f"- Observed coded values absent from the official guide: **{len(unexpected)}**.",
        f"- Invalid candidate-field formats: **{format_issue_count}**.",
        f"- Collision severity codes are exactly 1/2/3: **{severity_expected}**.",
        "",
        "## Annual rows",
        "",
    ]
    lines.extend(f"- {year}: {row_counts[year]:,}" for year in YEARS)
    lines.extend(
        [
            "",
            "## Interpretation",
            "",
            "All seven annual files use the same current 44-column schema because they were extracted from one current official complete file. This confirms structural consistency, but it does not by itself prove that every field is suitable for prediction.",
            "",
            "The complete file provides both historic variables and DfT-harmonized current replacements for junction detail, pedestrian crossing and carriageway hazards. The current harmonized variables should normally be preferred; the historic versions are retained only for traceability and sensitivity checks.",
            "",
            "`collision_severity` is consistently coded as 1=Fatal, 2=Serious and 3=Slight. `enhanced_severity_collision`, `collision_injury_based`, `collision_adjusted_severity_serious`, `collision_adjusted_severity_slight`, and `number_of_casualties` are outcome-related fields. Their final exclusion is a D3 leakage-audit decision, not a D2 schema decision.",
            "",
            "Category numbers are labels, not continuous magnitudes. Unknown or missing categories must remain explicit during later feature engineering rather than being silently treated as ordinary ordered numbers.",
        ]
    )
    if unexpected:
        lines.extend(
            [
                "",
                "## Unexpected codes requiring review",
                "",
            ]
        )
        lines.extend(
            f"- {row['year']} {row['field_name']}: {row['observed_code']}"
            for row in unexpected[:50]
        )

    (LOG_DIR / "d2_summary.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    dictionary, _ = read_official_dictionary()
    headers = get_headers()
    schema_rows, code_rows, row_counts = audit_data(headers, dictionary)
    mapping_rows = build_mapping(headers, dictionary)

    write_csv(LOG_DIR / "d2_schema_by_year.csv", schema_rows)
    write_csv(LOG_DIR / "d2_code_by_year.csv", code_rows)
    write_csv(LOG_DIR / "d2_field_mapping.csv", mapping_rows)
    write_summary(headers, schema_rows, code_rows, row_counts)

    print("D2 audit completed.")
    print(f"  years: {YEARS[0]}-{YEARS[-1]}")
    print(f"  fields per year: {len(headers[YEARS[0]])}")
    print(f"  total rows: {sum(row_counts.values()):,}")
    print(f"  outputs: {LOG_DIR}")


if __name__ == "__main__":
    main()
