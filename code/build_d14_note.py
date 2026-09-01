from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


PROJECT = Path(r"C:\Users\Administrator\Desktop\STATS19论文")
OUTPUT = PROJECT / "manuscript" / "D14问题确认与补充说明.docx"


# Resolved preset: standard_business_brief, with Microsoft YaHei as the East
# Asian fallback so Chinese text remains readable in Word and LibreOffice.
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "172B4D"
MUTED = "667085"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
CAUTION = "7A5A00"
RISK = "9B1C1C"
GREEN = "1F6B4F"
WHITE = "FFFFFF"


def set_run_font(run, *, latin="Calibri", east="Microsoft YaHei", size=11,
                 color="172B4D", bold=None, italic=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, *, latin="Calibri", east="Microsoft YaHei", size=11,
                   color=INK, bold=None):
    style.font.name = latin
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_border(cell, *, color="D0D5DD", size="4", val="single"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), val)
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, margins=CELL_MARGINS):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "bottom", "start", "end"):
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(margins[side]))
        element.set(qn("w:type"), "dxa")


def apply_table_geometry(table, widths):
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {CONTENT_WIDTH_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for idx, width in enumerate(widths):
        table.columns[idx].width = Twips(width)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Twips(widths[idx])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[idx]))
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)


def add_numbering(doc, *, left=720, hanging=360):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "decimal")
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "%1.")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    ppr.append(ind)
    for child in (start, fmt, text, suff, ppr):
        lvl.append(child)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_id = OxmlElement("w:abstractNumId")
    abs_id.set(qn("w:val"), str(abstract_id))
    num.append(abs_id)
    numbering.append(num)
    return num_id


def apply_number(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def add_text(paragraph, text, *, bold=False, italic=False, color=INK, size=11):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return run


def add_body(doc, text="", *, style="Normal", before=None, after=None,
             align=None, keep_next=False):
    p = doc.add_paragraph(style=style)
    if text:
        add_text(p, text)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    p.paragraph_format.keep_with_next = keep_next
    return p


def add_labeled_paragraph(doc, label, text, *, label_color=BLUE):
    p = doc.add_paragraph(style="Normal")
    add_text(p, label, bold=True, color=label_color)
    add_text(p, text)
    return p


def add_question(doc, num_id, title, status, paragraphs):
    p = doc.add_paragraph(style="Question")
    apply_number(p, num_id)
    add_text(p, title, bold=True, color=INK, size=12)
    add_text(p, f"  {status}", bold=True, color=(GREEN if "已确认" in status else CAUTION), size=9.5)
    p.paragraph_format.keep_with_next = True
    for label, text in paragraphs:
        add_labeled_paragraph(doc, label, text)


def style_table(table, *, header=True, font_size=9.2):
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell)
            if header and r_idx == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    set_run_font(run, size=font_size, color=INK, bold=(header and r_idx == 0))


def add_table(doc, headers, rows, widths, *, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.rows[0].cells
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        add_text(p, str(value), bold=True, color=INK, size=font_size)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cells[idx].text = ""
            p = cells[idx].paragraphs[0]
            add_text(p, str(value), size=font_size)
    apply_table_geometry(table, widths)
    style_table(table, font_size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_callout(doc, label, text, *, fill=CALLOUT, label_color=BLUE):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = ""
    set_cell_shading(cell, fill)
    set_cell_border(cell, color="C7D2E1", size="6")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    add_text(p, label, bold=True, color=label_color, size=10.5)
    add_text(p, text, size=10.5)
    apply_table_geometry(table, [CONTENT_WIDTH_DXA])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    add_text(p, text, bold=True, color=(BLUE if level < 3 else DARK_BLUE), size=(16 if level == 1 else 13 if level == 2 else 12))
    p.paragraph_format.keep_with_next = True
    return p


def configure_styles(doc):
    normal = doc.styles["Normal"]
    set_style_font(normal, size=11, color=INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for level, size, color, before, after in (
        (1, 16, BLUE, 16, 8),
        (2, 13, BLUE, 12, 6),
        (3, 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[f"Heading {level}"]
        set_style_font(style, size=size, color=color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Question" not in [s.name for s in doc.styles]:
        question = doc.styles.add_style("Question", WD_STYLE_TYPE.PARAGRAPH)
    else:
        question = doc.styles["Question"]
    set_style_font(question, size=12, color=INK, bold=True)
    question.paragraph_format.space_before = Pt(8)
    question.paragraph_format.space_after = Pt(4)
    question.paragraph_format.line_spacing = 1.10
    question.paragraph_format.keep_with_next = True


def configure_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    add_text(hp, "D14 | STATS19 事故严重度预测", size=9, color=MUTED, bold=True)
    add_text(hp, "    结果核查与论文写作补充说明", size=9, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    add_text(fp, "D14问题确认与补充说明  |  第 ", size=9, color=MUTED)
    add_page_field(fp)
    add_text(fp, " 页", size=9, color=MUTED)


def build_document():
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    question_num = add_numbering(doc)
    checklist_num = add_numbering(doc)

    # Title block (memo_masthead pattern).
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    add_text(p, "D14问题确认与补充说明", bold=True, color="000000", size=23)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    add_text(p, "STATS19 事故严重度预测项目 | SHAP、敏感性分析与收尾核查", color="4B5563", size=13)

    for label, value in (
        ("文档用途", "核对 D14 的17项方法、结果和写作问题，并形成正文同步依据"),
        ("核查日期", "2026年8月31日"),
        ("D14状态", "已完成；存在一项预定的地区留出分析未执行，已登记偏差"),
        ("判断原则", "事实、合理解释与论文处理建议分开表述；不把 SHAP 贡献解释为因果效应"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.10
        add_text(p, f"{label}: ", bold=True, color=INK, size=10.5)
        add_text(p, value, color=INK, size=10.5)

    add_callout(
        doc,
        "结论先行：",
        "D14 的 SHAP 稳定性分析、排除2020敏感性、错误结构审计和独立自检均已完成。正式稿仍须同步三项关键边界：A8 所说的“相关特征组及方向分布”本轮没有实施，不能写成已完成；special_conditions_at_site 的排名突升应作为潜在编码/分布漂移信号解释；地区泛化没有实证检验，不能声称跨辖区迁移性能。",
    )

    add_heading(doc, "一、H3/SHAP 方法", 1)
    add_body(doc, "本节回答 SHAP 的模型对象、聚合口径、抽样和排名解释。已完成的分析使用固定模型和预先冻结的抽样协议；第8项是唯一需要修改正文承诺的 SHAP 事项。")

    add_question(doc, question_num, "SHAP 用的是哪个模型？", "已确认", [
        ("核实结果：", "H3 使用5个冻结的随机参考 LightGBM 模型，随机种子为1103、2207、3301、4409、5501。5个模型均使用 D10 确定的 C03-1200 超参数，但分别独立训练。每个比较中，使用同一个已经拟合的模型分别解释其自身随机内部测试样本和同一批2024样本。"),
        ("边界：", "严格时间模型的2024 SHAP 结果仅作为补充描述，不参与 H3 的主相关分析。H3 解释的是“固定拟合模型面对内部测试分布与2024分布时，其特征重要性排序如何变化”，不是重新比较不同模型，也不是用一个 C03-1200 模型替代5个种子模型。"),
        ("正文写法：", "方法中明确写“对每个冻结的随机参考 LightGBM，固定模型参数和模型本身，仅改变解释样本分布；比较内部测试集与2024测试集的 SHAP 排名”。"),
    ])

    add_question(doc, question_num, "SHAP 是三类分别计算，还是三类合并计算？", "已确认", [
        ("核实结果：", "先对 Slight、Serious、Fatal 三个输出分别计算全部解释记录上的平均绝对 SHAP 值；主分析的总体重要性是三个类别特异性平均绝对 SHAP 的等权算术平均。SHAP 解释的是 LightGBM 多分类原始得分，类别顺序固定为 Slight、Serious、Fatal。"),
        ("正文写法：", "建议使用“class-specific mean absolute SHAP”与“equal-class mean importance”两个术语，并避免写成把三类记录简单混成一个标签后再计算 SHAP。"),
    ])

    add_question(doc, question_num, "Spearman 是在哪个排序上计算的？", "已确认", [
        ("核实结果：", "对17个冻结输入特征分别得到平均绝对 SHAP 重要性，按重要性排序；再比较同一种子模型在随机内部测试样本和2024样本上的两个17特征排序，计算 Spearman 等级相关系数。它不是对逐条记录的 SHAP 值与标签计算相关，也不是对 one-hot 展开后的全部列单独做一套未说明的排序。"),
        ("正文写法：", "写清“Spearman 相关在17个原始输入特征的平均绝对 SHAP 排名之间计算”。"),
    ])

    add_question(doc, question_num, "10,000条分层抽样中 Fatal 有多少？够不够稳定？", "已核实，需克制解释", [
        ("精确数量：", "共同的2024解释样本包含 Slight 7,516 条、Serious 2,335 条、Fatal 149 条；5个随机内部解释样本各包含 Slight 7,796 条、Serious 2,059 条、Fatal 145 条。每组均为10,000条，抽样不放回并按真实严重度分层、按比例整数分配。"),
        ("重要澄清：", "Fatal 输出的平均绝对 SHAP 是在全部10,000条解释记录上计算的，不是只在149条或145条真实 Fatal 记录上计算。因此 Fatal 数量影响的是样本组成和其代表性，不等于 Fatal 输出 SHAP 的唯一计算样本量。"),
        ("判断：", "149/145条对固定模型的探索性解释是可运行的，但不足以支持对致命类别特征贡献作高精度、普遍化的推断。Bootstrap 只能量化冻结模型和解释样本条件下的不确定性，不能补足真实 Fatal 总量、训练不确定性或跨年份稳定性。正文应报告分层样本数，并把 Fatal 特异性排序定位为描述性/探索性结果。"),
    ])

    add_question(doc, question_num, "Bootstrap 2000次是否重跑 TreeSHAP？", "已确认", [
        ("核实结果：", "不是。TreeSHAP 只对每个固定解释样本计算一次；随后在记录层面、按真实严重度分层有放回重采样2,000次，基于已经保存的绝对 SHAP 值重新计算平均重要性和 Spearman 排名相关，再取百分位区间。内部样本和2024样本独立重采样；共同的2024样本沿用同一组冻结抽样规则。"),
        ("正文写法：", "明确这是“record-level stratified bootstrap of precomputed SHAP values”，而不是2,000次重新拟合模型或重新运行 TreeSHAP。"),
    ])

    add_question(doc, question_num, "ρ=0.606、SD和范围的报告口径是什么？", "已确认", [
        ("核实结果：", "总体 rho=0.605882，跨5个种子的样本标准差为0.027260，范围为0.571078-0.632353。这三个数是5个种子各自得到的点估计 rho 的均值、标准差和范围，不是把5个种子的特征重要性合并成一个 pooled ranking。"),
        ("区间口径：", "每个种子的95%区间来自该种子内部与2024比较的记录层面 Bootstrap；它不是“5个种子均值”的置信区间。5个随机划分存在重叠，跨种子汇总应称为描述性汇总，不当作5项独立研究。"),
    ])

    add_table(doc,
              ["种子", "比较对象", "rho", "95% Bootstrap区间", "解释"],
              [
                  ("1103", "内部 vs 2024", "0.6152", "[0.5956, 0.6201]", "固定该种子模型"),
                  ("2207", "内部 vs 2024", "0.5833", "[0.5735, 0.5907]", "固定该种子模型"),
                  ("3301", "内部 vs 2024", "0.5711", "[0.5662, 0.5956]", "固定该种子模型"),
                  ("4409", "内部 vs 2024", "0.6324", "[0.6103, 0.6691]", "固定该种子模型"),
                  ("5501", "内部 vs 2024", "0.6275", "[0.6029, 0.6299]", "固定该种子模型"),
              ],
              [900, 1900, 1100, 2200, 3260],
              font_size=9.0)
    add_body(doc, "表1  五个随机参考模型的总体 SHAP 排名相关（区间为各模型内部的记录层面百分位区间）。", style="Caption", after=8)

    add_question(doc, question_num, "Top 特征是什么？哪些特征发生了位置交换？", "已核实，需避免机制化解读", [
        ("总体排序：", "内部测试样本的平均前5位为 speed_limit（1.0）、road_type（2.8）、hour（3.4）、weather_conditions（4.2）和 second_road_class（4.8）；2024样本的平均前5位为 special_conditions_at_site（1.0）、speed_limit（2.0）、hour（4.0）、second_road_class（4.8）和 road_type（5.0）。括号内为跨5个种子的平均名次。"),
        ("主要变化：", "special_conditions_at_site 从平均第17位升至第1位，5个种子方向一致；junction_detail 从10.6升至6.2；first_road_class 从8.6降至12.0。road_type、junction_control、pedestrian_crossing 和 weather_conditions 也出现一致或近一致的名次后移。"),
        ("关键风险：", "项目数据审计显示，special_conditions_at_site 的“Data missing or out of range”编码占比从2023年的约2.99%升至2024年的约58.75%。这不能单独证明编码漂移造成了 SHAP 变化，但足以说明排名突升可能与字段编码/分布变化相伴。正文应将其写成待核查的数据分布漂移信号，不得写成道路条件对事故严重度的真实因果机制。"),
    ])

    add_table(doc,
              ["特征", "内部平均名次", "2024平均名次", "名次变化（2024-内部）", "5个种子方向一致性"],
              [
                  ("special_conditions_at_site", "17.0", "1.0", "-16.0", "5/5 上升"),
                  ("junction_detail", "10.6", "6.2", "-4.4", "5/5 上升"),
                  ("first_road_class", "8.6", "12.0", "+3.4", "5/5 后移"),
                  ("road_type", "2.8", "5.0", "+2.2", "5/5 后移"),
                  ("speed_limit", "1.0", "2.0", "+1.0", "5/5 后移"),
              ],
              [3000, 1250, 1250, 1850, 2010],
              font_size=8.9)
    add_body(doc, "表2  总体等权类别 SHAP 排名的主要变化。名次变化为负表示在2024排序中上升。", style="Caption", after=8)

    add_question(doc, question_num, "大纲所说的“相关特征组及方向分布”做了吗？", "未实施，必须改正文承诺", [
        ("事实：", "没有。D14 已实施的是绝对 SHAP 重要性、17特征排序及其 Spearman 稳定性；相关特征组识别和有符号 SHAP 方向分布没有生成正式结果。"),
        ("处理建议：", "当前最稳妥的做法是把大纲中的“同时展示相关特征组及方向分布”删掉，正文只承诺已经完成的绝对 SHAP 排名稳定性。如果坚持保留，必须另行冻结补充协议，预先规定相关性阈值、分析层级（原始字段还是 one-hot 列）、类别范围、方向统计方式和多重比较处理，然后把它作为新增的补充分析，不能事后悄悄补写。"),
    ])

    add_heading(doc, "二、排除2020敏感性分析", 1)
    add_body(doc, "该分析回答的是训练期组成变化下主结论是否保持，不是对疫情年份的因果识别。")

    add_question(doc, question_num, "排除2020后是重新训练，还是只改变评估集？", "已确认", [
        ("核实结果：", "两种模型都重新训练。训练年份为2018、2019、2021、2022，训练样本447,262条；验证集仍为2023年104,258条；测试集仍为同一批2024年100,927条记录。预处理器和类别权重在缩减后的训练集上重新拟合。"),
        ("固定项：", "LightGBM 仍固定使用 C03-1200，不重新调参。这样比较的是“去掉2020并按同一模型协议重训”的敏感性，而不是偷偷改变模型选择。"),
    ])

    add_question(doc, question_num, "排除2020后 Macro-F1 增益变大说明什么？", "已确认，不能作因果解释", [
        ("结果：", "主分析中 LightGBM 相对加权 Logistic 的 Macro-F1 增益为+0.013896；排除2020后为+0.020623，95%区间为[+0.017398, +0.023892]。两种设置方向一致，但增益幅度增加约0.006728，区间为[+0.003924, +0.009490]。"),
        ("正确解释：", "这说明结论方向对训练期组成具有一定敏感性，效应幅度会随训练年份构成变化；它不能证明2020年本身造成了变化，因为训练样本量、年份构成和重新计算的类别权重同时发生了改变。"),
    ])

    add_table(doc,
              ["指标", "主分析差值", "排除2020差值", "排除2020区间", "判定"],
              [
                  ("Macro-F1（LGBM-Logistic）", "+0.013896", "+0.020623", "[+0.017398, +0.023892]", "方向不变，幅度增加"),
                  ("Fatal recall（LGBM-Logistic）", "-0.100533", "-0.091877", "[-0.106525, -0.076565]", "仍明显下降"),
              ],
              [2500, 1450, 1450, 2100, 1860],
              font_size=8.9)
    add_body(doc, "表3  主分析与排除2020敏感性分析的关键差值。", style="Caption", after=8)

    add_question(doc, question_num, "排除2020后 H2 联合判据是否仍不通过？", "已确认", [
        ("核实结果：", "仍不通过。Macro-F1 增益满足预设的实际增益规则，但 Fatal recall 差值为-0.091877，95%区间为[-0.106525, -0.076565]，表现为清晰的严重类别召回下降。"),
        ("论文结论：", "敏感性分析支持主结论的定性判断：LightGBM 并非在所有安全相关指标上都优于加权 Logistic，存在 Macro-F1 改善与 Fatal recall 恶化的指标权衡。不能仅凭 Macro-F1 宣称统一增量价值。"),
    ])

    add_heading(doc, "三、错误结构与安全指标", 1)
    add_body(doc, "这里的错误结构来自2024测试集全体100,927条记录。Fatal总数为1,502条，计数均以真实类别为条件汇总。")

    add_question(doc, question_num, "Fatal→Slight 与 Fatal→Serious 的完整结构是多少？", "已核实", [
        ("核实结果：", "加权 Logistic 将952条 Fatal 误判为 Slight、223条误判为 Serious，正确识别327条；LightGBM 将897条误判为 Slight、429条误判为 Serious，正确识别176条。三类计数分别加总为1,502条。"),
        ("解释边界：", "LightGBM 减少了 Fatal→Slight 的远距离低估，但同时增加了 Fatal→Serious 的近距离低估，且正确 Fatal 数从327降至176。因此不能只引用952降到897，就说严重误判整体得到改善。"),
    ])

    add_table(doc,
              ["模型", "Fatal→Slight", "Fatal→Serious", "正确预测Fatal", "Fatal recall"],
              [
                  ("加权 Logistic", "952（63.38%）", "223（14.85%）", "327（21.77%）", "0.217710"),
                  ("加权 LightGBM", "897（59.72%）", "429（28.56%）", "176（11.72%）", "0.117177"),
              ],
              [2000, 1800, 1800, 1900, 1860],
              font_size=9.0)
    add_body(doc, "表4  真实 Fatal 样本的预测去向。百分比的分母为1,502条真实 Fatal 记录。", style="Caption", after=8)

    add_question(doc, question_num, "LightGBM 在2024上的 Fatal precision 是多少？", "已确认", [
        ("数值：", "Fatal precision=0.041519，约4.15%。LightGBM 在2024年共预测4,239条为Fatal，其中真正Fatal为176条，其余为非Fatal误报。"),
        ("论文用途：", "该数值适合放在完整指标表或混淆矩阵旁作为补充说明，用来呈现 Fatal recall 与 precision 的权衡。它不应被单独解释为模型“安全”或“不安全”的总体判定。"),
    ])

    add_question(doc, question_num, "误判数据进入哪些图表？", "建议已确定", [
        ("主文建议：", "主文放加权 Logistic 与 LightGBM 的配对3×3行归一化混淆矩阵，并在格内或旁注同时给出计数；主结果表列出 Macro-F1、QWK、等级 MAE、各类召回、Fatal precision/recall 和代价指标。"),
        ("文字讨论：", "在讨论中用一段话报告 Fatal→Slight 与 Fatal→Serious 的数量及方向，解释“远距离低估减少但正确 Fatal 也减少”的结构。不要另做一个只展示952与897的装饰性柱状图。"),
        ("附录：", "高置信度个体错误只作为确定性示例放附录，并明确“不代表总体错误分布”；不能把少量示例当成机制证据。"),
    ])

    add_heading(doc, "四、地区留出偏差与空间泛化", 1)
    add_body(doc, "地区留出原计划存在于大纲中，但在模型结果已知前没有冻结具体 police_force 组合。为保持审计诚实，D14 将其登记为未执行的计划分析，而不是事后挑选一个区域补跑。")

    add_question(doc, question_num, "deviation log 是否写清未做地区留出的原因和影响？", "已确认", [
        ("核实结果：", "已写清四层信息：原计划是什么；何时发现未冻结具体组合；证据是什么（D2 的 police_force 角色为 not_yet_decided、D3 决策仍为 PENDING，冻结配置中没有具体区域）；为什么现在不能再称为预设分析（结果已知后选择会变成 post-hoc）；以及采取的纠正措施和受影响的结论。"),
        ("写作边界：", "论文中应主动披露这是一项未执行的计划分析，不把它包装成“无结果的空间验证”。如以后补做，必须标为 post-hoc exploratory，并单独冻结新协议。"),
    ])

    add_question(doc, question_num, "局限性是否明确写了空间泛化未验证？", "已确认项目边界，正文需同步", [
        ("核实结果：", "D14 局限性文件已经明确：本研究未对按 police_force 划分的空间泛化进行实证评价，因此不能声称模型性能可跨辖区迁移。"),
        ("正文建议句：", "可在局限性中写：“本研究未对按警察辖区划分的空间泛化进行实证检验，因而结果不支持跨辖区迁移性能结论；后续研究需在建模前预先冻结地区留出方案，并在独立辖区上验证。”"),
    ])

    add_heading(doc, "五、工程与复现核查", 1)
    add_body(doc, "三套独立自检脚本均已通过。它们验证的是协议、数据形状、结果对齐和文件哈希等可复现性条件，不等同于证明研究结论具有外部普适性。")

    add_question(doc, question_num, "三套独立测试分别测了什么？", "已确认", [
        ("SHAP 自检：", "核验冻结协议、确定性分层抽样、SHAP 维度与加和性、排名和区间是否由保存的结果重算得到，以及输出文件哈希是否一致。"),
        ("排除2020自检：", "核验训练/验证/测试年份和样本量、模型固定项、预测记录对齐、Bootstrap数组维度、模型参数和输出文件哈希。"),
        ("收尾自检：", "核验错误结构表、配对结果、确定性的附录示例、地区留出偏差记录、报告边界以及归档文件哈希。"),
    ])

    add_table(doc,
              ["自检项目", "一句话核验范围"],
              [
                  ("SHAP test", "协议、抽样、TreeSHAP形状/加和性、排名重算、Bootstrap维度与哈希"),
                  ("Exclude-2020 test", "年份与样本量、固定模型、预测对齐、Bootstrap维度与哈希"),
                  ("Closeout test", "错误表、配对结局、附录示例、地区偏差、边界和哈希"),
              ],
              [2300, 7060],
              font_size=9.2)
    add_body(doc, "表5  D14 三套独立自检的核验范围。", style="Caption", after=8)

    add_heading(doc, "六、正文同步清单", 1)
    add_body(doc, "在进入 D15 软件打包和论文撰写前，建议按下面顺序同步正文；其中前3项属于必须改动。")
    checklist = [
        "H3方法写明5个随机参考 LightGBM、同模型内部/2024比较、三类等权平均绝对 SHAP、17特征排名和 Bootstrap 口径，并报告各组 Fatal 数量。",
        "删除或改写“相关特征组及方向分布”的未实现承诺；若未来新增该分析，先冻结独立协议再执行。",
        "在结果和讨论中对 special_conditions_at_site 的排名突升加数据分布漂移警示，避免把 SHAP 排名写成因果机制。",
        "排除2020敏感性写明“重新训练两模型、2024测试记录不变、LightGBM不重新调参”，并说明方向不变但增益幅度变化。",
        "主文使用配对混淆矩阵和完整指标表，完整呈现 Fatal→Slight、Fatal→Serious、Fatal precision 与 Fatal recall 的权衡。",
        "局限性明确写空间泛化未验证，删除或改写任何跨 police_force 迁移性能暗示。",
        "保留代码版本、随机种子、线程限制、输入/输出哈希和三套自检记录，作为复现材料的一部分。",
    ]
    for item in checklist:
        p = doc.add_paragraph(style="Normal")
        apply_number(p, checklist_num)
        add_text(p, item)

    add_heading(doc, "七、核查依据", 1)
    add_body(doc, "本说明基于项目中以下冻结协议、结果文件和自检记录整理。路径以项目根目录 C:\\Users\\Administrator\\Desktop\\STATS19论文 为基准。")
    sources = [
        "config/d14_shap_protocol.json；results/d14/d14_explanation_samples.csv.gz；results/d14/d14_shap_importance.csv；results/d14/d14_rank_stability.csv",
        "results/d14/d14_rank_stability_summary.csv；results/d14/d14_error_structure.csv；results/d14/d14_paired_error_outcomes.csv",
        "results/d14/exclude2020/d14_exclude2020_validation_metrics.csv；results/d14/exclude2020/d14_exclude2020_test_metrics.csv；logs/d14_exclude2020_checkpoint.md",
        "logs/d14_protocol_deviations.json；logs/d14_limitations_and_boundaries.md；logs/d14_final_checkpoint.md；logs/d14_final_summary.json",
        "code/test_d14_shap.py；code/test_d14_exclude2020.py；code/test_d14_closeout.py（均已通过独立断言）",
    ]
    for item in sources:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        add_text(p, "• ", bold=True, color=BLUE)
        add_text(p, item, size=10)

    add_callout(
        doc,
        "D14结论：",
        "D14 可以关闭并进入 D15。关闭的含义是已有分析和审计结果已归档，不是所有原大纲承诺都已实现；A8 的未实施项、地区留出偏差和编码分布漂移风险必须在正式稿中如实处理。",
        fill="EEF7F2",
        label_color=GREEN,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
